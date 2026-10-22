from pathlib import Path
import shutil
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parent
PAGES = []

def page(title, *items):
    PAGES.append((title, items))

def p(text): return ('p', text)
def label(text): return ('label', text)
def code(text): return ('code', text.strip())
def table(headers, rows): return ('table', (headers, rows))

page('Tutorial de la base de datos del hospital',
    label('Guía desde cero para estudiantes y docentes'),
    p('Aprende a abrir el proyecto, ver sus tablas, ejecutar consultas, practicar cambios, crear un diagrama y hacer un respaldo. No necesitas saber programar para comenzar.'),
    p('Este manual corresponde al proyecto hospital_universitario: 36 tablas, 6 vistas, 6 procedimientos y 23 triggers. Usa datos ficticios y está pensado para una entrega universitaria.'),
    label('Qué necesitas'),
    p('Una computadora con Windows, MySQL Server, MySQL Workbench y los archivos de la carpeta propuesta_hospital. Si ya puedes conectarte al servidor, ve a la página 4.'),
    label('Elige por dónde empezar'),
    table(['Objetivo', 'Páginas'], [
        ['Ver el proyecto por primera vez', '2 y 4 a 8'],
        ['Instalar los programas', '3'],
        ['Practicar sin cambiar la entrega', '9 a 20'],
        ['Guardar un respaldo y recuperarlo', '21 y 22'],
        ['Presentar el proyecto al profesor', '23'],
        ['Editar un dato usando el ratón', '24'],
        ['Resolver errores y entender términos', '25 a 27']]),
    label('Regla para todos los ejercicios'),
    p('Lee primero el paso. Después selecciona y ejecuta solamente el bloque SQL indicado. Los cambios de los ejercicios se hacen en hospital_practica. Las instrucciones para crear las bases se ejecutan completas una sola vez.'),
    p('Versión de servidor utilizada para validar el proyecto original: MySQL 8.0.45. Requisito del SQL: MySQL 8.0.16 o posterior. Esta guía muestra menús de Workbench 8.0; algunas etiquetas pueden variar.'),
    p('Guarda la contraseña del servidor y la ubicación de tus archivos. No necesitas entregar tu contraseña al profesor.'))

page('1 Entender qué estás abriendo',
    p('Una base de datos organiza información en tablas. Imagina una tabla como una hoja con columnas y filas: cada fila de pacientes representa una persona ficticia; sus columnas guardan nombre, teléfono y otros datos.'),
    table(['Elemento', 'Qué hace'], [
        ['MySQL Server', 'Guarda los datos y ejecuta las instrucciones SQL. Debe estar encendido.'],
        ['MySQL Workbench', 'Es la interfaz que usas para conectarte, consultar y enviar cambios al servidor.'],
        ['Archivo .sql', 'Es texto con instrucciones. Abrirlo no instala ni cambia la base.'],
        ['Archivo .docx o .pdf', 'Es documentación para leer. No crea tablas.'],
        ['Archivo .svg', 'Es un dibujo que puedes abrir en un navegador y ampliar.'],
        ['Archivo .mwb', 'Es un modelo editable de Workbench, con diagramas.']]),
    label('Recorrido rápido para ver el proyecto'),
    p('1. Extrae hospital_universitario.zip: clic derecho sobre el ZIP, Extraer todo y Extraer. Trabaja con los archivos extraídos.'),
    p('2. Abre MySQL Workbench y entra a una conexión que funcione. La página 4 explica cómo crearla.'),
    p('3. Si hospital_universitario todavía no existe, abre hospital_completo.sql y ejecútalo completo como se explica en la página 6.'),
    p('4. En Schemas, refresca el panel y abre hospital_universitario → Tables → pacientes. Elige Select Rows – Limit… con clic derecho.'),
    p('5. Abre diagramas/vision_modulos.svg para ubicar las áreas del hospital y diagramas/er_completo.svg para ver sus relaciones.'),
    label('Lo que se guarda en cada lugar'),
    p('Los datos viven en el servidor. El archivo SQL contiene instrucciones y el Word explica cómo usarlas. Copiar la carpeta del proyecto a otra computadora no copia los cambios que ya hiciste en el servidor; para eso necesitas un respaldo.'))

page('2 Instalar MySQL y Workbench',
    p('Si ya tienes MySQL Server y Workbench funcionando, salta esta página. Workbench por sí solo no es el servidor.'),
    label('Ruta para seguir la versión 8 punto 0 del proyecto'),
    p('1. Entra a https://dev.mysql.com/downloads/installer/ y descarga MySQL Installer Community para Windows. La serie de ese asistente es 8.0. No es la versión más nueva de MySQL.'),
    p('2. Ejecuta el instalador. En Setup Type, elige Custom. En la selección de productos agrega MySQL Server 8.0 y MySQL Workbench. Si aparece una actualización de mantenimiento de la serie 8.0, anota la versión instalada.'),
    p('3. En Check Requirements, resuelve los requisitos indicados por el asistente. Ejecuta la instalación y espera a que termine.'),
    p('4. En la configuración del servidor, selecciona Development Computer o Development según la etiqueta. Usa TCP/IP y anota el puerto, normalmente 3306.'),
    p('5. Define una contraseña para root. Guárdala. Sigue la configuración como servicio de Windows, aplica los cambios y comprueba que terminó correctamente.'),
    p('6. Abre Workbench desde Inicio. Si no aparece, confirma que lo seleccionaste en el instalador.'),
    label('Si instalas MySQL 8 punto 4'),
    p('Descarga Server 8.4 LTS desde https://dev.mysql.com/downloads/mysql/ y Workbench desde https://dev.mysql.com/downloads/workbench/. El MSI de Server 8.4 usa MySQL Configurator; Workbench se instala aparte. Configura Development, TCP/IP, contraseña de root y servicio de Windows; luego aplica la configuración.'),
    p('Los scripts originales se probaron en Server 8.0.45. Workbench se desarrolló y probó con Server 8.0; su manual advierte que algunas funciones gráficas pueden fallar con 8.4 o versiones posteriores. No interpretes esa advertencia de conexión como prueba de que el SQL falló.'),
    label('Si ya existe un servidor'),
    p('Usa el puerto y usuario de ese servidor. Si el puerto 3306 está ocupado, identifica qué instancia lo utiliza antes de crear otra. El puerto 33060 es de otro protocolo y no se usa en esta conexión Standard TCP/IP.'),
    p('Fuentes oficiales de instalación y compatibilidad: página 27.'))

page('3 Crear y probar la conexión',
    p('Una conexión le dice a Workbench dónde está el servidor y con qué usuario entrar. El nombre de la conexión es solo una etiqueta.'),
    p('1. En la pantalla inicial de Workbench, localiza MySQL Connections y pulsa el signo +.'),
    p('2. Completa los campos usando los datos de tu instalación:'),
    table(['Campo', 'Valor para un servidor local típico'], [
        ['Connection Name', 'Hospital local'],
        ['Connection Method', 'Standard (TCP/IP)'],
        ['Hostname', '127.0.0.1 o localhost'],
        ['Port', '3306, o el puerto que anotaste'],
        ['Username', 'root, o el usuario que te proporcionaron'],
        ['Default Schema', 'Déjalo vacío en esta primera conexión']]),
    p('3. Pulsa Test Connection. Escribe la contraseña que definiste al instalar el servidor.'),
    p('4. Si la prueba funciona, pulsa OK. Luego abre la conexión desde la pantalla inicial.'),
    p('5. En una pestaña SQL nueva, pega y ejecuta esta consulta:'),
    code('''SELECT VERSION() AS version_servidor,
       @@version_comment AS descripcion;'''),
    p('Resultado esperado: una fila con la versión y descripción del servidor. Si aparece MariaDB, no es el motor MySQL validado en este proyecto; la etiqueta MySQL de XAMPP no basta para identificarlo.'),
    label('Si no puedes entrar'),
    p('Connection refused o error 2003: comprueba que el servicio del servidor está iniciado y que el puerto sea correcto. Access denied o error 1045: revisa usuario y contraseña. Workbench debe conectarse antes de que puedas ejecutar SQL.'),
    p('No hay una contraseña universal del proyecto. Los registros ficticios de la tabla usuarios tampoco son cuentas para conectarse a MySQL.'))

page('4 Ubicar los botones y ejecutar SQL',
    p('Una vez conectada la sesión, identifica estas zonas de Workbench:'),
    table(['Zona', 'Para qué sirve'], [
        ['Schemas a la izquierda', 'Muestra bases, tablas, vistas y rutinas.'],
        ['Pestaña SQL en el centro', 'Aquí escribes o pegas instrucciones.'],
        ['Barra con botones de rayo', 'Envía SQL al servidor.'],
        ['Result Grid debajo del editor', 'Muestra filas devueltas por una consulta.'],
        ['Action Output en la parte inferior', 'Muestra éxito, errores y mensajes de ejecución.']]),
    label('Para cada bloque de este manual'),
    p('1. Abre una pestaña de consulta con el botón SQL + o File → New Query Tab. Pega el bloque sin las marcas de formato del documento.'),
    p('2. Selecciona con el ratón todas las líneas del bloque que quieres ejecutar. Incluye los puntos y coma.'),
    p('3. Usa Query → Execute (All or Selection), o Ctrl + Shift + Enter. El primer botón de rayo hace esta misma operación: si no hay texto seleccionado, ejecuta toda la pestaña.'),
    p('4. Revisa Result Grid y Action Output. Si un paso da error, resuélvelo antes de continuar con el siguiente.'),
    p('Activa Query → Stop Script Execution On Errors para que un error detenga el bloque. Si se detiene dentro de una transacción de práctica, ejecuta ROLLBACK antes de continuar.'),
    label('Instrucción actual frente a bloque completo'),
    p('Ctrl + Enter ejecuta la instrucción donde está el cursor. Es útil para una consulta aislada. Los ejemplos con SET y SELECT o con START TRANSACTION y ROLLBACK necesitan que selecciones el bloque completo.'),
    label('Guardar y ejecutar son acciones diferentes'),
    p('Ctrl + S guarda el texto en un archivo .sql. No ejecuta ese texto ni modifica el servidor. Ejecutar SQL envía las instrucciones al servidor; el resultado depende de cada instrucción y de la transacción.'),
    p('No ejecutes todo este tutorial como si fuera un instalador. Incluye lecturas, cambios, operaciones que solo deben hacerse una vez y errores intencionales.'))

page('5 Crear la base del proyecto',
    p('Haz esto una sola vez por servidor cuando hospital_universitario todavía no exista. Si ya aparece, pasa a la siguiente página.'),
    p('1. En Workbench conectado, abre File → Open SQL Script, o Ctrl + Shift + O.'),
    p('2. Selecciona propuesta_hospital/hospital_completo.sql de la carpeta extraída.'),
    p('3. Comprueba que la pestaña corresponde a ese archivo. No selecciones un fragmento: necesitas ejecutar el archivo completo.'),
    p('4. Usa Query → Execute (All or Selection). Espera a que termine y revisa Action Output. Las instrucciones DELIMITER pertenecen al script de las rutinas; no las borres.'),
    p('5. En Schemas, haz clic derecho y elige Refresh All. Despliega hospital_universitario.'),
    p('6. En una pestaña nueva, ejecuta:'),
    code('''USE hospital_universitario;
SELECT DATABASE() AS base_actual;
SELECT COUNT(*) AS pacientes_iniciales
FROM pacientes;'''),
    p('Resultado inicial esperado: hospital_universitario y 5 pacientes ficticios. Si ya modificaste datos, el conteo puede ser diferente.'),
    label('Elige una forma de instalación'),
    p('hospital_completo.sql reúne estructura, reglas y datos. También existen 01_esquema.sql, 02_reglas.sql y 03_datos_demo.sql para instalarlos por separado en ese orden. Usa una sola opción; ejecutar ambas vuelve a intentar crear los mismos objetos y datos.'),
    label('Si aparece un error'),
    p('Error 1007: esa base ya existe. Detente; volver a ejecutar el archivo no actualiza ni guarda una instalación existente. Error 1050: una tabla ya existe, posiblemente por instalación parcial. Conserva la base y revisa el primer error de Action Output antes de decidir cómo corregirla.'),
    p('Este instalador no borra una base anterior. Los cinco archivos .ibd del prototipo son evidencia; no se abren como scripts ni se cargan automáticamente.'))

page('6 Ver las tablas y comprobar la instalación',
    label('Ver datos usando el ratón'),
    p('1. En Schemas, abre hospital_universitario → Tables. Si no lo ves, usa Refresh All.'),
    p('2. Haz clic derecho en pacientes y elige Select Rows – Limit… El número del límite depende de la versión y preferencias de Workbench.'),
    p('3. Mira los nombres de las columnas y las filas ficticias de Result Grid. La pestaña SQL muestra la consulta que produjo ese resultado.'),
    label('Ver datos usando SQL'),
    code('''USE hospital_universitario;
SELECT paciente_id, expediente_clinico, nombre,
       apellido_paterno, fecha_nacimiento, telefono
FROM pacientes
ORDER BY paciente_id;'''),
    p('Esperas 5 filas en la carga inicial. ORDER BY ordena el resultado; no cambia el orden físico ni los datos guardados.'),
    label('Contar tablas y vistas'),
    code('''SELECT table_type AS tipo, COUNT(*) AS cantidad
FROM information_schema.tables
WHERE table_schema = DATABASE()
GROUP BY table_type;'''),
    p('Esperas BASE TABLE = 36 y VIEW = 6, es decir, 42 objetos entre ambos tipos. Una vista es una consulta guardada; no cuenta como tabla física del proyecto.'),
    label('Ver las columnas y la definición'),
    code('''DESCRIBE pacientes;
SHOW CREATE TABLE pacientes;'''),
    p('DESCRIBE muestra campos, tipos y claves. SHOW CREATE TABLE muestra el SQL completo de esa tabla. Si el texto se corta en la cuadrícula, usa clic derecho sobre la celda → Open Value in Editor para leerlo.'))

page('7 Abrir y crear los diagramas',
    label('Ver el dibujo ya entregado'),
    p('1. En el Explorador de archivos abre propuesta_hospital/diagramas.'),
    p('2. Abre vision_modulos.svg con Edge, Chrome u otro navegador para ubicar las áreas del hospital. Abre er_completo.svg para ver tablas, campos y relaciones; usa zoom. Los archivos er_ por módulo permiten leer áreas más pequeñas.'),
    p('3. Una línea conecta campos relacionados. El identificador de un paciente permite asociar citas, consultas y hospitalizaciones sin repetir toda su información en cada tabla.'),
    label('Crear tu propio diagrama editable en Workbench'),
    p('1. En Workbench conectado, abre Database → Reverse Engineer.'),
    p('2. Selecciona tu conexión y avanza con Next. Proporciona la contraseña si la pide.'),
    p('3. Selecciona solamente hospital_universitario. No agregues otras bases del servidor.'),
    p('4. Avanza hasta Select Objects to Reverse Engineer. Mantén las tablas y la opción Place imported objects on a diagram.'),
    p('5. Pulsa Execute, revisa el resultado y termina con Finish. Si hay varias páginas del modelo, abre la pestaña EER Diagram.'),
    p('6. Ordena las cajas para leerlas. Guarda el modelo con File → Save como hospital_modelo.mwb.'),
    label('Qué archivo consultar según la duda'),
    table(['Archivo del proyecto', 'Qué explica'], [
        ['MODELO.md', 'Alcance, relaciones, reglas y normalización.'],
        ['DICCIONARIO.md', 'Campos, tipos y claves de cada tabla.'],
        ['REVISION.md', 'Problemas del prototipo y mejoras realizadas.'],
        ['reporte_hospital.pdf', 'Informe académico completo.'],
        ['04_consultas.sql', '20 consultas para mostrar el funcionamiento.']]),
    p('Guardar o mover cajas en el modelo .mwb no cambia automáticamente el servidor. Cuando cambies la estructura mediante SQL, genera de nuevo el modelo para que el dibujo corresponda a la base actual.'))

page('8 Crear una base para practicar',
    p('Antes de modificar datos, crea hospital_practica. Es una instalación nueva con los mismos datos ficticios iniciales; no es un respaldo de los cambios que hiciste en hospital_universitario.'),
    p('1. Abre File → Open SQL Script y selecciona tutorial/hospital_practica.sql.'),
    p('2. Ejecuta ese archivo completo una sola vez. No selecciones solo una parte.'),
    p('3. Revisa Action Output y refresca Schemas. Ahora deben aparecer hospital_universitario y hospital_practica si instalaste ambas.'),
    p('4. Abre una pestaña SQL nueva para los ejercicios y ejecuta:'),
    code('''USE hospital_practica;
SELECT DATABASE() AS base_actual,
       VERSION() AS version_mysql;
SHOW FULL TABLES;'''),
    p('Esperas base_actual = hospital_practica y 42 objetos de tabla o vista. Los siguientes ejercicios parten de la carga inicial y se realizan en orden.'),
    label('Comprueba la base antes de cada cambio'),
    p('USE selecciona la base de esta sesión. Si abres otra conexión, vuelve a escribir USE hospital_practica. Como verificación rápida, ejecuta SELECT DATABASE(); y lee el nombre.'),
    label('Si hospital_practica ya existe'),
    p('No vuelvas a ejecutar su instalador para reiniciarla. Puedes continuar con ella teniendo en cuenta los cambios que ya hiciste. Para otro inicio limpio, copia el archivo y reemplaza todas las apariciones del nombre hospital_practica por uno nuevo, por ejemplo hospital_practica_2, usando Find and Replace del editor. Comprueba el texto antes de ejecutar y usa ese nuevo nombre en tus consultas.'),
    p('No cambies solo el nombre en CREATE DATABASE: las instrucciones USE del mismo archivo también deben apuntar al nombre nuevo. Conserva intactos los archivos originales del proyecto.'))

page('9 Consultar datos sin modificarlos',
    p('SELECT lee filas. Puedes ejecutar estas consultas varias veces: no cambian registros. Selecciona un bloque a la vez y revisa sus resultados.'),
    label('Mostrar columnas y ordenar'),
    code('''USE hospital_practica;
SELECT paciente_id, nombre, apellido_paterno, telefono
FROM pacientes
ORDER BY paciente_id;'''),
    p('SELECT indica qué columnas mostrar, FROM dice de qué tabla y ORDER BY ordena la lista. El punto y coma termina la instrucción.'),
    label('Filtrar con una condición'),
    code('''SELECT paciente_id, nombre, sexo, activo
FROM pacientes
WHERE activo = 1 AND sexo = 'F'
ORDER BY paciente_id;'''),
    p('En los datos iniciales aparecen las pacientes con IDs 1 y 4. AND pide que se cumplan las dos condiciones. El texto va entre comillas simples.'),
    label('Buscar parte de un texto'),
    code('''SELECT paciente_id, nombre
FROM pacientes
WHERE nombre LIKE '%Alex%';'''),
    p('Esperas al paciente con ID 3. El símbolo % permite que haya más caracteres antes o después de Alex.'),
    label('Contar en lugar de listar'),
    code('''SELECT COUNT(*) AS total_pacientes
FROM pacientes;'''),
    p('Esperas 5 antes de agregar otro paciente. COUNT(*) cuenta filas. AS da un nombre fácil de leer a una columna del resultado.'),
    p('Una consulta sin resultados puede ser correcta: significa que ninguna fila cumplió la condición. Una fila de resultado con 0 en un conteo no es lo mismo que un error.'))

page('10 Consultar dos tablas relacionadas',
    p('Los datos de un paciente y los de sus contactos de emergencia viven en tablas distintas. JOIN permite mostrarlos juntos.'),
    code('''USE hospital_practica;
SELECT p.paciente_id, p.nombre AS paciente,
       ce.nombre AS contacto,
       ce.telefono AS telefono_contacto
FROM pacientes AS p
LEFT JOIN contactos_emergencia AS ce
  ON ce.paciente_id = p.paciente_id
 AND ce.prioridad = 1
ORDER BY p.paciente_id;'''),
    p('Resultado inicial: 5 filas, una por paciente, con su contacto de prioridad 1 cuando existe.'),
    label('Cómo leer esta consulta'),
    p('p es un nombre corto para pacientes y ce para contactos_emergencia. p.nombre significa la columna nombre de pacientes.'),
    p('ON une el contacto con el paciente que tiene el mismo paciente_id. La condición prioridad = 1 elige el contacto principal.'),
    p('LEFT JOIN mantiene al paciente aunque no tenga ese contacto. En ese caso, las columnas del contacto aparecerán como NULL.'),
    label('Por qué importan las claves'),
    p('paciente_id es la clave primaria de pacientes: identifica una fila sin ambigüedad. En contactos_emergencia es una clave foránea: señala a un paciente existente.'),
    p('Dos personas pueden tener el mismo nombre. Por eso la relación usa un identificador y no el nombre de la persona. Tampoco debes escribir listas de contactos en una sola celda.'),
    label('Practica una consulta del proyecto'),
    p('Abre 04_consultas.sql y busca una pregunta que te interese. Su instrucción USE apunta a hospital_universitario. Para probarla en tu copia, abre otra pestaña, escribe USE hospital_practica y pega solamente la consulta elegida, sin copiar el USE anterior.'),
    p('Los ejercicios de transacción comentados al final de ese archivo son ejemplos adicionales. No necesitas ejecutarlos para ver sus 20 consultas de lectura.'))

page('11 Agregar un paciente ficticio',
    p('INSERT agrega una fila. Usaremos un expediente exclusivo del tutorial para reconocer el registro después.'),
    p('1. Comprueba que la base es hospital_practica. Ejecuta este bloque una sola vez en la misma conexión que usarás en las dos páginas siguientes:'),
    code('''USE hospital_practica;
INSERT INTO pacientes (
  expediente_clinico, nombre, apellido_paterno,
  apellido_materno, fecha_nacimiento, sexo, telefono
)
VALUES (
  'EXP-2026-0031', 'Nicolás', 'Herrera',
  'Ríos', '2000-01-15', 'NO_ESPECIFICADO',
  '5512349087'
);'''),
    p('2. Comprueba en Action Output que INSERT terminó sin error. Solo si tuvo éxito, ejecuta este segundo bloque en la misma conexión:'),
    code('''SET @paciente_tutorial = LAST_INSERT_ID();
SELECT paciente_id, expediente_clinico, nombre, telefono
FROM pacientes
WHERE paciente_id = @paciente_tutorial;'''),
    p('3. Anota el paciente_id devuelto. Debe aparecer Nicolás Herrera con su expediente y teléfono ficticios.'),
    label('Lo que hizo el bloque'),
    p('AUTO_INCREMENT genera el identificador. LAST_INSERT_ID() recupera el que acaba de generarse en esta conexión; no supongas que siempre será 6. La variable @paciente_tutorial guarda ese ID durante la sesión.'),
    p('Las columnas que omitimos usan su valor predeterminado o permiten NULL según la definición de la tabla. Las fechas se escriben año-mes-día.'),
    label('Si sale el error 1062'),
    p('El expediente ya existe. Detente: no ejecutes el SET ni las modificaciones siguientes basándote en un INSERT que falló. Para repetir el ejercicio, elige otro expediente exclusivo, cambia también los filtros posteriores y vuelve a insertar.'),
    p('Si cerraste la conexión, la variable se pierde. Puedes recuperar el registro con SELECT paciente_id FROM pacientes WHERE expediente_clinico = \'EXP-2026-0031\'; y establecer SET @paciente_tutorial = el_numero_real; reemplazando el_numero_real por el ID que leíste.'))

page('12 Cambiar un dato y deshacerlo',
    p('UPDATE cambia filas. Una transacción permite revisar un cambio antes de confirmarlo. Este ejercicio usa el paciente agregado en la página anterior.'),
    p('1. En la misma conexión, selecciona y ejecuta el bloque completo:'),
    code('''USE hospital_practica;
START TRANSACTION;

SELECT paciente_id, telefono AS antes
FROM pacientes
WHERE paciente_id = @paciente_tutorial;

UPDATE pacientes
SET telefono = '5587654301'
WHERE paciente_id = @paciente_tutorial;

SELECT paciente_id, telefono AS durante
FROM pacientes
WHERE paciente_id = @paciente_tutorial;

ROLLBACK;

SELECT paciente_id, telefono AS despues
FROM pacientes
WHERE paciente_id = @paciente_tutorial;'''),
    p('2. Revisa las tres cuadrículas: antes y después deben mostrar 5512349087; durante debe mostrar 5587654301.'),
    label('Cómo guardar un cambio'),
    p('START TRANSACTION abre una transacción. ROLLBACK deshace sus cambios pendientes; COMMIT los confirma. Si quieres guardar el teléfono nuevo en este ejercicio, repite el bloque sustituyendo ROLLBACK por COMMIT.'),
    p('Con autocommit activo, una instrucción de modificación fuera de una transacción se confirma automáticamente. No confíes en Ctrl + Z para deshacer un cambio del servidor: ese atajo solo afecta al editor.'),
    label('Evita cambiar todas las filas'),
    p('Antes de UPDATE, usa SELECT con el mismo WHERE y confirma qué fila apunta. Aquí WHERE usa una clave primaria. Sin WHERE, el cambio podría afectar a todos los pacientes.'),
    p('No uses este patrón para los procedimientos de negocio ni para CREATE o ALTER. Esos casos tienen reglas diferentes que se explican más adelante.'))

page('13 Eliminar solo el registro de práctica',
    p('DELETE elimina filas. Este ejemplo borra únicamente el paciente que acabas de crear; todavía no tiene citas, consultas u otros registros asociados.'),
    p('1. Ejecuta primero esta lectura y confirma que muestra Nicolás Herrera:'),
    code('''USE hospital_practica;
SELECT paciente_id, expediente_clinico, nombre
FROM pacientes
WHERE paciente_id = @paciente_tutorial
  AND expediente_clinico = 'EXP-2026-0031';'''),
    p('2. Si aparece exactamente el registro correcto, ejecuta el bloque siguiente:'),
    code('''DELETE FROM pacientes
WHERE paciente_id = @paciente_tutorial
  AND expediente_clinico = 'EXP-2026-0031';

SELECT ROW_COUNT() AS filas_eliminadas;

SELECT paciente_id, nombre
FROM pacientes
WHERE paciente_id = @paciente_tutorial;'''),
    p('Resultado esperado: filas_eliminadas = 1 y una última cuadrícula sin filas. Si aparece 0, revisa el ID, el expediente y si el registro ya se eliminó.'),
    label('Por qué se agrega el expediente al filtro'),
    p('El identificador y el expediente exclusivo permiten comprobar juntos que estamos borrando el registro de este ejercicio. Si usaste otro expediente, cambia ambos filtros para que coincidan con el que insertaste.'),
    label('Por qué puede fallar eliminar un paciente'),
    p('La base protege las relaciones. Un paciente con consultas o citas asociadas normalmente no se puede borrar directamente: el error 1451 indica que hay registros que dependen de él.'),
    p('No borres registros relacionados para saltarte esa protección. Para conservar historia, muchas entidades permiten marcar activo = 0 mediante el flujo que corresponda. Este tutorial usa DELETE solo sobre el registro ficticio recién creado y sin dependencias.'),
    p('Si Workbench devuelve 1175, consulta la página 25. Conserva un filtro con la clave primaria en lugar de desactivar globalmente Safe Updates.'))

page('14 Agendar una cita con un procedimiento',
    p('Un procedimiento agrupa instrucciones para una operación. sp_agendar_cita revisa el intervalo y evita solapamientos del médico o del paciente.'),
    label('Antes de empezar'),
    p('Termina cualquier transacción pendiente con COMMIT para confirmar o ROLLBACK para deshacer lo que corresponda. Los seis procedimientos de negocio de este proyecto abren y confirman sus propias transacciones. No los envuelvas en otra transacción esperando deshacerlos después.'),
    p('1. En hospital_practica con sus datos iniciales, ejecuta una sola vez:'),
    code('''USE hospital_practica;
SET @cita_tutorial = NULL;
CALL sp_agendar_cita(
  4, 1,
  '2026-10-15 15:00:00',
  '2026-10-15 15:30:00',
  'Control de seguimiento',
  @cita_tutorial
);
SELECT @cita_tutorial AS cita_creada;
SELECT cita_id, paciente_id, medico_id,
       inicio, fin, estado
FROM citas
WHERE cita_id = @cita_tutorial;'''),
    p('2. Esperas una cita para el paciente 4, médico 1 y el intervalo indicado. El último parámetro devuelve su identificador.'),
    label('Cómo interpretar los argumentos'),
    p('4 es paciente_id y 1 es medico_id. Las dos fechas son inicio y fin. El texto es el motivo y @cita_tutorial recibe el resultado de salida. No reemplaces números sin verificar primero que existen en sus tablas.'),
    p('Las fechas fijas permiten reproducir la demostración; no representan el día actual. Al repetir exactamente esta llamada, debe rechazarse el solapamiento de la cita ya creada. Eso demuestra una regla, no una falla de instalación.'),
    p('Los parámetros y reglas de los seis procedimientos están en 02_reglas.sql. Esta guía muestra tres operaciones completas.'))

page('15 Consultar hospitalizaciones y camillas',
    p('Una camilla puede estar operativa y, al mismo tiempo, ocupada. Su disponibilidad combina el estado operativo con una asignación abierta.'),
    code('''USE hospital_practica;
SELECT camilla_id, codigo, estado_operativo,
       disponibilidad
FROM v_camillas_estado
ORDER BY camilla_id;

SELECT disponibilidad, COUNT(*) AS cantidad
FROM v_camillas_estado
GROUP BY disponibilidad;

SELECT hospitalizacion_id, paciente_id,
       fecha_ingreso, fecha_alta
FROM hospitalizaciones
ORDER BY hospitalizacion_id;'''),
    p('En la carga inicial hay 1 camilla ocupada, 4 disponibles y 1 en mantenimiento. La hospitalización 1 tiene alta registrada y la 2 sigue abierta.'),
    label('Qué significa NULL'),
    p('NULL indica que no hay un valor registrado. Cuando fecha_alta aparece como NULL, el alta todavía no está registrada; no es una fecha cero ni una cadena vacía. Para buscarlo usa IS NULL:'),
    code('''SELECT hospitalizacion_id, paciente_id, fecha_ingreso
FROM hospitalizaciones
WHERE fecha_alta IS NULL;'''),
    label('Cómo se hacen los cambios de hospitalización'),
    p('sp_asignar_camilla asigna una camilla y, cuando procede, cierra la asignación anterior para conservar el traslado. sp_alta_hospitalaria registra el alta y cierra la asignación abierta. Consulta sus argumentos en 02_reglas.sql antes de practicarlos.'),
    p('No cambies manualmente un texto de ocupación para liberar una camilla. La ocupación se deriva de las asignaciones. v_camillas_estado refleja asignaciones con fin NULL; no reconstruye la ocupación de cualquier fecha histórica que escribas.'),
    p('Una vista facilita leer valores calculados. En este proyecto, consulta las vistas y realiza los cambios mediante las tablas o los procedimientos adecuados.'))

page('16 Practicar una entrega de farmacia',
    p('Este ejercicio registra una entrega parcial ficticia. Hazlo una sola vez para comparar los resultados indicados. Mantén los registros de inventario en su orden cronológico.'),
    label('Leer antes de entregar'),
    code('''USE hospital_practica;
SELECT receta_id, renglon, cantidad_prescrita,
       cantidad_entregada, cantidad_pendiente
FROM v_recetas_pendientes
WHERE receta_id = 2 AND renglon = 1;
SELECT lote_id, stock
FROM v_stock_lotes WHERE lote_id = 3;'''),
    p('Valores iniciales: 4 pendientes en ese renglón y 40 unidades en el lote 3.'),
    label('Registrar la operación'),
    code('''SET @dispensacion_tutorial = NULL;
CALL sp_dispensar(
  2, 1, 3, 1, 6,
  '2026-10-10 10:00:00',
  @dispensacion_tutorial
);
SELECT @dispensacion_tutorial AS entrega_creada;'''),
    p('Argumentos: receta 2, renglón 1, lote 3, cantidad 1, personal de farmacia 6 y fecha. El procedimiento registra la entrega y el movimiento de inventario en una operación.'),
    label('Comprobar después'),
    p('Repite solamente las dos consultas de lectura de arriba. Esperas 1 entregada, 3 pendientes y stock = 39.'),
    p('Si vuelves a ejecutar CALL, puedes registrar otra entrega real de práctica y bajar otra vez el stock mientras haya cantidad pendiente. No es un botón para refrescar ni una operación que se pueda repetir sin consecuencias.'),
    p('El stock se calcula sumando movimientos. No hagas UPDATE para escribir una existencia nueva. Los movimientos y detalles de entrega son inmutables; las correcciones de inventario usan un nuevo ajuste justificado. Los datos de sustancias y cantidades del proyecto son ficticios.'))

page('17 Practicar un pago parcial',
    p('El saldo de una factura se calcula a partir de sus conceptos y pagos. Usa una factura emitida y ejecuta la llamada una sola vez.'),
    label('Leer el saldo inicial'),
    code('''USE hospital_practica;
SELECT factura_id, folio, estado, total, pagado, saldo
FROM v_facturas_saldo
WHERE factura_id = 1;'''),
    p('Valores iniciales de factura 1: total = 850.00, pagado = 500.00 y saldo = 350.00.'),
    label('Registrar un abono ficticio'),
    code('''SET @pago_tutorial = NULL;
CALL sp_registrar_pago(
  1, '2026-10-10 10:30:00', 25.00,
  'EFECTIVO', 'EFE-202610-000025',
  @pago_tutorial
);
SELECT @pago_tutorial AS pago_creado;'''),
    label('Leer de nuevo'),
    code('''SELECT factura_id, total, pagado, saldo
FROM v_facturas_saldo
WHERE factura_id = 1;'''),
    p('Esperas pagado = 525.00 y saldo = 325.00. El total permanece en 850.00.'),
    label('Qué recordar'),
    p('Repetir CALL agrega otro pago; la referencia no funciona como un bloqueo contra duplicados. Ejecuta la operación solo cuando quieres registrar un abono nuevo.'),
    p('Una factura ABIERTA admite conceptos; primero se emite para recibir pagos. La factura 4 del ejemplo todavía está ABIERTA. Los pagos registrados no se editan ni se eliminan en este alcance académico.'),
    p('No escribas el saldo manualmente en una tabla: la vista lo calcula. Este módulo es una demostración de cobros y no un sistema de facturación fiscal.'))

page('18 Comprobar reglas con errores esperados',
    p('Ahora ejecutarás tres instrucciones incorrectas a propósito. Selecciona y ejecuta una sola instrucción a la vez. Debe aparecer un error y no debe agregarse la fila.'),
    label('Un código único repetido'),
    code('''USE hospital_practica;
INSERT INTO departamentos (codigo, nombre)
VALUES ('DEP-001', 'Consulta externa');'''),
    p('Esperas error 1062 Duplicate entry: ya existe ese código. UNIQUE protege valores que no deben repetirse.'),
    label('Un valor fuera de la regla'),
    code('''INSERT INTO contactos_emergencia (
  paciente_id, nombre, parentesco, telefono, prioridad
)
VALUES (1, 'Teresa Herrera Ríos', 'Hermana',
        '5512349088', 0);'''),
    p('Esperas error 3819 Check constraint: prioridad debe ser mayor que 0. CHECK valida una condición del registro.'),
    label('Una relación hacia un paciente inexistente'),
    code('''INSERT INTO contactos_emergencia (
  paciente_id, nombre, parentesco, telefono, prioridad
)
VALUES (999999, 'Teresa Herrera Ríos', 'Hermana',
        '5512349088', 1);'''),
    p('Esperas error 1452 Foreign key constraint: el paciente 999999 no existe en la carga inicial.'),
    label('Cómo explicarlo al profesor'),
    p('Estos errores demuestran que el servidor protege reglas del modelo. No debes desactivar CHECK o claves foráneas para conseguir que entren datos incorrectos.'),
    p('Las reglas entre filas, como citas que se solapan o entregas que superan una receta, también necesitan procedimientos y transacciones. Una futura aplicación debe usar esos procedimientos y restringir las escrituras directas que eludan sus comprobaciones.'))

page('19 Modificar la estructura de una tabla',
    p('Modificar datos cambia filas. Modificar estructura cambia columnas, tipos, índices o relaciones. Para aprender ALTER, usaremos una tabla nueva de prueba.'),
    p('1. En hospital_practica ejecuta una sola vez:'),
    code('''USE hospital_practica;
CREATE TABLE prueba_estructura (
  id INT UNSIGNED AUTO_INCREMENT PRIMARY KEY,
  nombre VARCHAR(100) NOT NULL
);
DESCRIBE prueba_estructura;'''),
    p('Esperas dos columnas: id y nombre. INT guarda números enteros; VARCHAR(100) admite texto de hasta 100 caracteres; NOT NULL exige un valor.'),
    p('2. Ejecuta una sola vez:'),
    code('''ALTER TABLE prueba_estructura
  ADD COLUMN observaciones VARCHAR(200) NULL;
DESCRIBE prueba_estructura;'''),
    p('Esperas una tercera columna observaciones. NULL permite dejarla sin valor registrado. La copia de práctica ahora tiene 37 tablas físicas porque agregaste una tabla de ejercicio.'),
    label('Un cambio de estructura no se deshace igual'),
    p('CREATE y ALTER producen confirmaciones implícitas en MySQL. No prometas recuperar un cambio de estructura con ROLLBACK. Haz respaldo y revisa qué se alterará antes de aplicarlo a una tabla del proyecto.'),
    label('Si cambias realmente el proyecto'),
    p('Actualiza los scripts de creación y reglas afectados, el diccionario, el reporte y los diagramas. Repite la validación con los archivos nuevos. La evidencia 82 de 82 del proyecto original no prueba automáticamente una versión que tú modificaste.'),
    p('Error 1050 al crear la tabla o 1060 al agregar la columna: probablemente ya ejecutaste ese paso. Usa DESCRIBE para comprobarlo antes de volver a intentarlo.'))

page('20 Respaldar la base con Workbench',
    p('Un respaldo SQL conserva estructura y datos del servidor en un archivo. Copiar hospital_completo.sql conserva la instalación inicial, pero no los cambios posteriores.'),
    p('1. Conéctate al servidor donde está la base que quieres respaldar.'),
    p('2. Abre Server → Data Export. Marca hospital_universitario, o hospital_practica si quieres guardar tus ejercicios.'),
    p('3. Marca todas las tablas de esa base. Selecciona Dump Structure and Data.'),
    p('4. En las opciones, activa Dump Stored Procedures and Functions y Dump Triggers. Mantén los eventos si existieran; este proyecto no necesita eventos. Revisa que las vistas queden incluidas en la exportación del esquema.'),
    p('5. Selecciona Export to Self-Contained File: se generará un solo archivo SQL. Elige una carpeta que puedas encontrar.'),
    p('6. Activa Include Create Schema si aparece. Así el archivo puede incluir la creación y selección de la base.'),
    p('7. Usa un nombre que identifique la base y el momento, por ejemplo hospital_universitario_2026_10_10.sql. No reemplaces el único respaldo que tengas.'),
    p('8. Pulsa Start Export. Revisa Export Progress y comprueba que terminó sin errores. Confirma en el Explorador que el archivo existe y no está vacío.'),
    label('Qué comprobar en el archivo'),
    p('Ábrelo para revisar el texto sin ejecutarlo. Debes encontrar definiciones de tablas, datos, vistas y rutinas. La presencia de CREATE DATABASE, USE o DROP TABLE influye en dónde se restaurará y qué reemplazará.'),
    label('Guardar el trabajo del equipo'),
    p('Conserva el respaldo junto a una copia fechada de los scripts, documentos y modelo .mwb. Indica qué cambió y quién lo cambió. Si todos editan una base compartida, acuerden una sola versión antes de exportar.'),
    p('Un respaldo se comprueba de verdad restaurándolo en otro entorno y consultando sus datos. Si la exportación marca un problema de compatibilidad de herramientas, revisa las versiones; no ignores el error ni entregues el archivo como respaldo exitoso.'),
    p('Fuente oficial y opciones de exportación: página 27.'))

page('21 Recuperar un respaldo y compartir',
    label('Restaurar en otra computadora o instancia'),
    p('Para principiantes, utiliza una instancia donde no exista una base importante con el mismo nombre. No pruebes la restauración sobre la única copia de tu trabajo.'),
    p('1. Guarda el archivo de respaldo SQL en la computadora destino. Instala y configura MySQL y Workbench si hace falta.'),
    p('2. Conéctate al servidor destino y abre Server → Data Import.'),
    p('3. Elige Import from Self-Contained File y selecciona el respaldo.'),
    p('4. Antes de importar, abre el archivo como texto y revisa los nombres de CREATE DATABASE y USE, y si tiene DROP TABLE o DROP DATABASE. Si ya existe una base con ese nombre, detente y elige otro entorno para probar.'),
    p('5. Si el respaldo incluye la creación del esquema, úsala. Si no la incluye, crea una base vacía con el nombre correcto y selecciona Default Target Schema según lo que contenga el archivo.'),
    p('6. Pulsa Start Import y revisa Import Progress. Debe terminar sin errores.'),
    p('7. Refresca Schemas. Comprueba tablas, vistas, rutinas y algunos datos de la base restaurada.'),
    label('El nombre destino requiere revisar el SQL'),
    p('Elegir otra base en Default Target Schema no redirige necesariamente un archivo que ya contiene CREATE DATABASE o USE. Un respaldo con DROP TABLE puede reemplazar tablas existentes. No asumas que el asistente lo convertirá en una copia independiente.'),
    label('Qué entregar o pasar a un compañero'),
    table(['Si necesita', 'Comparte'], [
        ['La instalación ficticia inicial', 'hospital_completo.sql y documentos del proyecto.'],
        ['Tus datos y cambios actuales', 'Un respaldo SQL completo verificado.'],
        ['Seguir los ejercicios', 'Este manual y tutorial/hospital_practica.sql.'],
        ['Editar el dibujo', 'Tu modelo .mwb o las fuentes Mermaid .mmd.']]),
    p('Cada compañero debe utilizar sus propias credenciales del servidor. No escriban contraseñas en el reporte o en archivos que vayan a entregar.'))

page('22 Mostrar el proyecto en cinco minutos',
    p('Este recorrido sirve para un estudiante que expone y para un docente que quiere revisar el resultado sin conocer Workbench. Solo usa lecturas.'),
    p('1. Abre reporte_hospital.pdf y diagramas/vision_modulos.svg. Explica que el alcance es un sistema académico de atención, hospitalización, farmacia, procedimientos y cobros.'),
    p('2. En Workbench entra a la conexión. Ejecuta:'),
    code('''USE hospital_universitario;
SELECT DATABASE() AS base;
SELECT COUNT(*) AS pacientes FROM pacientes;
SELECT table_type AS tipo, COUNT(*) AS cantidad
FROM information_schema.tables
WHERE table_schema = DATABASE()
GROUP BY table_type;'''),
    p('En la carga original muestra 5 pacientes, 36 tablas y 6 vistas. Si el equipo ya cambió datos, explícalo y muestra su respaldo actualizado.'),
    p('3. Muestra atención clínica y ocupación:'),
    code('''SELECT * FROM v_expediente_consultas
ORDER BY paciente_id, fecha;
SELECT codigo, disponibilidad
FROM v_camillas_estado
ORDER BY camilla_id;'''),
    p('4. Muestra los valores calculados de farmacia y cobros:'),
    code('''SELECT lote_id, stock FROM v_stock_lotes;
SELECT folio, total, pagado, saldo
FROM v_facturas_saldo;'''),
    p('5. Abre el ER y señala una clave primaria y una foránea. Explica con pacientes y contactos por qué se relacionan mediante paciente_id.'),
    label('Explicar la evidencia con precisión'),
    p('VALIDACION.json registra 82 de 82 comprobaciones aprobadas del proyecto original en MySQL 8.0.45. El tutorial pasó 28 de 28 comprobaciones de SQL en ese servidor; la evidencia está en tutorial/VALIDACION_TUTORIAL.json. Las pruebas SQL no son una prueba visual de la interfaz de Workbench.'),
    p('Aclara que los datos son ficticios, la autenticación todavía no está implementada y la facturación es académica. Las tablas cubren los procesos incluidos en MODELO.md; no pretenden representar cada sistema de un hospital real.'))

page('23 Cambiar un dato usando el ratón',
    p('También puedes editar una celda de una tabla directamente en Result Grid. Practica solo en hospital_practica y con un dato sencillo, como el teléfono ficticio.'),
    p('1. Ejecuta esta consulta para mostrar un único paciente con su clave primaria:'),
    code('''SELECT paciente_id, nombre, telefono
FROM hospital_practica.pacientes
WHERE paciente_id = 1;'''),
    p('2. Anota el teléfono que aparece antes de cambiarlo. Haz doble clic en la celda telefono y escribe 5598760123. No edites paciente_id.'),
    p('3. Pulsa Apply en la zona de resultados. Workbench abre una ventana para revisar el SQL que enviará al servidor.'),
    p('4. Lee el cambio antes de aprobarlo. Debe actualizar pacientes de hospital_practica, cambiar solo telefono y filtrar el paciente con ID 1. Si muestra otra tabla o varias filas, cancela y revisa el paso.'),
    p('5. En la ventana de revisión pulsa Apply. Comprueba que terminó correctamente y pulsa Finish.'),
    p('6. Vuelve a ejecutar la consulta de arriba o refresca la cuadrícula. El teléfono debe mostrar 5598760123.'),
    p('7. Para devolver el dato al valor original, edita la misma celda con el valor que anotaste y repite Apply, revisión, Apply y Finish.'),
    label('Cuándo no se puede editar la cuadrícula'),
    p('Si aparece Read Only, el resultado no es editable en ese contexto. Puede faltar la clave primaria, combinar varias tablas o no tener permisos de escritura. Usa una consulta de una sola tabla con su PK o el SQL del ejercicio correspondiente.'),
    label('Las reglas también se aplican aquí'),
    p('Apply envía cambios reales al servidor; no es solo guardar el texto. Los triggers y restricciones pueden rechazarlos. Evita editar pagos, movimientos de inventario o entregas: este proyecto los protege para conservar su historia.'),
    p('La ruta se basa en el manual oficial Result Grid. Si tu versión muestra una ventana distinta, lee el SQL y sus mensajes antes de confirmar.'))

page('24 Resolver los errores más comunes',
    table(['Mensaje o situación', 'Qué revisar y hacer'], [
        ['No aparece la base en Schemas', 'Refresh All. Comprueba la conexión y que la instalación terminó.'],
        ['2003 o Connection refused', 'Servidor iniciado, host y puerto correctos. Revisa el servicio de Windows.'],
        ['1045 Access denied', 'Usuario y contraseña del servidor. No uses usuarios ficticios de la tabla usuarios.'],
        ['1046 No database selected', 'Ejecuta USE hospital_practica o la base que necesitas.'],
        ['1049 Unknown database', 'Nombre mal escrito o base no creada en esta instancia.'],
        ['1007 o 1050 Already exists', 'No reinstales encima. Revisa si el paso ya se ejecutó o hubo instalación parcial.'],
        ['1062 Duplicate entry', 'Código o expediente único repetido. Comprueba el registro existente.'],
        ['1064 Syntax error', 'Revisa comillas simples, comas, nombres y punto y coma. En reglas ejecuta el script completo con DELIMITER.'],
        ['1146 Table does not exist', 'Base seleccionada, nombre de tabla y ejecución de 01_esquema.sql.'],
        ['1175 Safe update mode', 'Usa WHERE con la clave primaria. Si la variable no se reconoce, usa el ID numérico leído y añade el expediente para confirmar.'],
        ['1451 Cannot delete parent row', 'Existen registros relacionados. Conserva la historia y revisa el flujo correspondiente.'],
        ['1452 Cannot add child row', 'La clave foránea apunta a un registro inexistente. Consulta primero el ID correcto.'],
        ['3819 Check constraint violated', 'Un valor incumple una regla. Revisa el campo y la definición CHECK.'],
        ['1644 o mensaje del proyecto', 'Una rutina o trigger rechazó la operación. Lee el mensaje: puede indicar solapamiento, saldo, stock o estado inválido.'],
        ['No cambia nada o devuelve 0 filas', 'Revisa WHERE, variable, base, conexión y si el cambio ya se hizo. No siempre es un error.']]),
    p('Action Output conserva el mensaje de la operación. Para pedir ayuda comparte el texto exacto del primer error, el bloque que ejecutaste y la versión del servidor. Oculta contraseñas y cualquier dato real.'))

page('25 Glosario para seguir la clase',
    table(['Palabra', 'Significado en este proyecto'], [
        ['Base o esquema', 'Contenedor de tablas y otros objetos; en MySQL estos nombres se usan como equivalentes.'],
        ['Tabla', 'Conjunto de filas con las mismas columnas.'],
        ['Fila o registro', 'Una instancia, como un paciente ficticio.'],
        ['Columna o campo', 'Un atributo, como telefono.'],
        ['Tipo de dato', 'Qué valores admite una columna: texto, número, fecha y otros.'],
        ['SQL', 'Lenguaje para consultar y modificar bases relacionales.'],
        ['PK o clave primaria', 'Identificador único y obligatorio de una fila.'],
        ['FK o clave foránea', 'Referencia a una clave de otra tabla o de la misma tabla.'],
        ['UNIQUE', 'Evita repetir un valor o combinación de valores.'],
        ['CHECK', 'Exige que una condición del registro se cumpla.'],
        ['NULL', 'Ausencia de valor registrado. Se consulta con IS NULL.'],
        ['AUTO_INCREMENT', 'Genera identificadores numéricos al insertar.'],
        ['Vista', 'Consulta guardada que se lee como una tabla.'],
        ['Procedimiento', 'Rutina que se invoca con CALL y puede recibir y devolver parámetros.'],
        ['Trigger', 'Regla que se ejecuta automáticamente ante una operación en una tabla.'],
        ['Transacción', 'Grupo de cambios que se confirma con COMMIT o se deshace con ROLLBACK cuando sigue pendiente.'],
        ['Autocommit', 'Modo que confirma automáticamente instrucciones fuera de una transacción explícita.'],
        ['JOIN', 'Combina datos de tablas mediante una condición de relación.'],
        ['ER o EER', 'Diagrama de entidades, atributos y relaciones.'],
        ['Respaldo', 'Archivo que permite reconstruir estructura y datos de un momento.']]),
    label('Cuatro verbos que conviene distinguir'),
    p('Abrir muestra un archivo. Guardar conserva el texto o modelo. Ejecutar envía SQL al servidor. Confirmar con COMMIT hace permanentes los cambios pendientes de una transacción.'))

page('26 Fuentes oficiales y comprobación final',
    p('Los menús de este manual se contrastaron con documentación oficial de MySQL. Las páginas de descarga pueden mostrar versiones nuevas; anota siempre la versión que realmente instalaste.'),
    label('Descargas e instalación'),
    p('MySQL Installer 8.0: https://dev.mysql.com/downloads/installer/'),
    p('Productos y modo Custom: https://dev.mysql.com/doc/refman/8.0/en/mysql-installer-setup.html'),
    p('MySQL Server: https://dev.mysql.com/downloads/mysql/'),
    p('MySQL Workbench: https://dev.mysql.com/downloads/workbench/'),
    p('Server 8.4 en Windows: https://dev.mysql.com/doc/refman/8.4/en/windows-installation.html'),
    label('Uso de Workbench'),
    p('Conexión: https://dev.mysql.com/doc/workbench/en/wb-mysql-connections-new.html'),
    p('Barra SQL: https://dev.mysql.com/doc/workbench/en/wb-sql-editor-toolbar.html'),
    p('Menú Query: https://dev.mysql.com/doc/workbench/en/wb-sql-editor-main-menu.html'),
    p('Edición visual: https://dev.mysql.com/doc/workbench/en/wb-develop-sql-editor-results.html'),
    p('Atajos: https://dev.mysql.com/doc/workbench/en/wb-keys.html'),
    p('Schemas: https://dev.mysql.com/doc/workbench/en/wb-sql-editor-navigator.html'),
    p('Diagrama: https://dev.mysql.com/doc/workbench/en/wb-reverse-engineer-live.html'),
    p('Exportación e importación: https://dev.mysql.com/doc/workbench/en/wb-admin-export-import-management.html'),
    p('Compatibilidad de Workbench: https://dev.mysql.com/doc/workbench/en/'),
    label('Antes de entregar'),
    p('Comprueba que puedes conectarte, abrir las tablas y mostrar el diagrama. Exporta un respaldo con las rutinas y triggers. Entrega archivos que se correspondan con la misma versión del proyecto.'),
    p('Las instrucciones y datos de ejemplo se validan por separado del formato del Word y PDF. Si cambias fechas, IDs, scripts o datos, los resultados esperados de los ejercicios pueden cambiar.'))

def shade(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), color); tcpr.append(shd)

def add_table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    if headers[0] == 'Objetivo': widths = [Cm(11.8), Cm(5.6)]
    elif headers[0] == 'Archivo del proyecto': widths = [Cm(6.0), Cm(11.4)]
    elif headers[0] == 'Mensaje o situación': widths = [Cm(6.2), Cm(11.2)]
    else: widths = [Cm(5.1), Cm(12.3)]
    for col, width in zip(t.columns, widths): col.width = width
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]; cell.width = widths[i]
        cell.text = header; shade(cell, 'F2F2F2')
        for r in cell.paragraphs[0].runs: r.bold = True
    trpr = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values): cells[i].text = value; cells[i].width = widths[i]
    borders = OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el = OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    t._tbl.tblPr.append(borders)
    for row in t.rows:
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, val in [('top','70'),('bottom','70'),('left','100'),('right','100')]:
                el = OxmlElement('w:'+side); el.set(qn('w:w'),val); el.set(qn('w:type'),'dxa'); margins.append(el)
            tcpr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05
                for r in para.runs: r.font.size = Pt(10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def build():
    ROOT.mkdir(exist_ok=True)
    src = (PROJECT / 'hospital_completo.sql').read_text(encoding='utf-8')
    (ROOT / 'hospital_practica.sql').write_text('-- COPIA NUEVA PARA PRACTICAR. Ejecutar completa una sola vez.\n-- No es un respaldo de datos modificados.\n' + src.replace('hospital_universitario', 'hospital_practica'), encoding='utf-8')
    doc = Document()
    sec = doc.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.7); sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.8); sec.right_margin = Cm(1.8)
    sec.header_distance = Cm(.7); sec.footer_distance = Cm(.7)
    styles = doc.styles
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        styles[name].font.name = 'Calibri'; styles[name].font.color.rgb = RGBColor(0,0,0)
        for border in list(styles[name]._element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    normal = styles['Normal']; normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    styles['Title'].font.size = Pt(26); styles['Title'].font.bold = True
    styles['Heading 1'].font.size = Pt(19); styles['Heading 1'].font.bold = True
    styles['Heading 1'].paragraph_format.space_after = Pt(14)
    styles['Heading 2'].font.size = Pt(12); styles['Heading 2'].font.bold = True
    styles['Heading 2'].paragraph_format.space_before = Pt(8)
    styles['Heading 2'].paragraph_format.space_after = Pt(5)
    header = sec.header.paragraphs[0]; header.text = 'HOSPITAL UNIVERSITARIO   |   GUÍA DE USO'
    for r in header.runs: r.font.size = Pt(8); r.font.color.rgb = RGBColor(100,100,100)
    footer = sec.footer.paragraphs[0]; footer.alignment = 2
    r = footer.add_run('Página '); r.font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    md=[]
    for idx,(title,items) in enumerate(PAGES):
        if idx: doc.add_page_break()
        doc.add_paragraph(title, 'Title' if idx==0 else 'Heading 1')
        md += ['# '+title, '']
        for kind,value in items:
            if kind in ('p','label'):
                para = doc.add_paragraph(style='Heading 2' if kind=='label' else 'Normal')
                if idx == len(PAGES)-1 and 'https://' in value:
                    prefix, url = value.split('https://',1); url = 'https://' + url
                    para.add_run(prefix).font.size = Pt(9)
                    link = OxmlElement('w:hyperlink')
                    link.set(qn('r:id'), doc.part.relate_to(url,RT.HYPERLINK,is_external=True))
                    r = OxmlElement('w:r'); rpr = OxmlElement('w:rPr')
                    size = OxmlElement('w:sz'); size.set(qn('w:val'),'18'); rpr.append(size)
                    color = OxmlElement('w:color'); color.set(qn('w:val'),'1F4E79'); rpr.append(color)
                    r.append(rpr); text = OxmlElement('w:t'); text.text = url; r.append(text); link.append(r); para._p.append(link)
                else: para.add_run(value)
                md += [('## ' if kind=='label' else '')+value, '']
            elif kind=='code':
                para = doc.add_paragraph()
                para.paragraph_format.space_before = Pt(3); para.paragraph_format.space_after = Pt(9)
                para.paragraph_format.line_spacing = 1.04
                para.paragraph_format.keep_together = True
                r=para.add_run(value); r.font.name='Consolas'; r.font.size=Pt(9.3)
                md += ['```sql',value,'```','']
            elif kind=='table':
                headers,rows=value; add_table(doc,headers,rows)
                md += ['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']
                md += ['| '+' | '.join(row)+' |' for row in rows]; md += ['']
    doc.core_properties.title='Tutorial de la base de datos del hospital'
    doc.core_properties.subject='Guía para principiantes de MySQL y Workbench'
    doc.core_properties.author=''
    doc.save(ROOT / 'tutorial_hospital_principiantes.docx')
    (ROOT / 'TUTORIAL.md').write_text('\n'.join(md),encoding='utf-8')
    print('DOCX, Markdown y base de práctica creados; páginas previstas:',len(PAGES))

if __name__ == '__main__': build()

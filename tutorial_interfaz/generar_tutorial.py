"""Convierte TUTORIAL.md en un manual Word editable.

Requiere python-docx. La edición del texto se hace en TUTORIAL.md.
El PDF final se obtiene con un renderizador de Word o LibreOffice.
"""

from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "tutorial_interfaz_hospital.docx"
TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def hyperlink(paragraph, label, url):
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "174A70")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def inline(paragraph, text):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10.5)
        elif part.startswith("["):
            match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            hyperlink(paragraph, match.group(1), match.group(2))
        else:
            paragraph.add_run(part)


def remove_borders(paragraph_style):
    properties = paragraph_style.element.find(qn("w:pPr"))
    if properties is not None:
        for node in list(properties):
            if node.tag == qn("w:pBdr"):
                properties.remove(node)


def add_table(document, headers, rows):
    widths = [2.05, 4.75] if len(headers) == 2 else [6.8 / len(headers)] * len(headers)
    table = document.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    table_properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    table_properties.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for edge, value in (("top", "95"), ("bottom", "95"), ("left", "110"), ("right", "110")):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    table_properties.append(margins)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, values in enumerate([headers, *rows]):
        row = table.rows[0] if index == 0 else table.add_row()
        for col, text in enumerate(values):
            cell = row.cells[col]
            cell.width = Inches(widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EBEF" if index == 0 else ("F7F8FA" if index % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shade)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.line_spacing = 1.05
            if len(rows) <= 6 and index < len(rows):
                paragraph.paragraph_format.keep_with_next = True
            inline(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(11)
                if index == 0:
                    run.bold = True
        # Individual rows should remain readable, while long tables may span pages.
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.2
    spacer.add_run().font.size = Pt(2)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.72)
    section.left_margin = section.right_margin = Inches(0.85)
    section.footer_distance = Inches(0.3)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        remove_borders(style)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11.5)
    normal.paragraph_format.line_spacing = 1.11
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    title = doc.styles["Title"]
    title.font.size = Pt(24)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Hospital universitario  |  ").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "Cómo usar la aplicación del hospital desde cero"
    doc.core_properties.subject = "Manual de la aplicación Python con MySQL"
    doc.core_properties.author = "Equipo del proyecto"
    doc.core_properties.keywords = "hospital, Python, MySQL, tutorial, interfaz"
    doc.core_properties.comments = ""
    lines = (ROOT / "TUTORIAL.md").read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == "<!-- pagebreak -->":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
            picture_path = ROOT / match.group(2)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(picture_path), width=Inches(6.65))
            caption = doc.add_paragraph(match.group(1))
            caption.paragraph_format.space_after = Pt(9)
            for run in caption.runs:
                run.font.size = Pt(10)
            i += 1
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:], "Title")
            i += 1
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 1")
            i += 1
            continue
        if line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
                i += 1
            add_table(doc, table_lines[0], table_lines[2:])
            continue
        paragraph = doc.add_paragraph()
        if re.match(r"^\d+\. ", line):
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.first_line_indent = Inches(-0.2)
        if line.startswith("- "):
            paragraph.paragraph_format.left_indent = Inches(0.2)
            paragraph.paragraph_format.first_line_indent = Inches(-0.2)
            line = "• " + line[2:]
        inline(paragraph, line)
        # Keep a short table introduction with its header and first rows.
        following = next((x.strip() for x in lines[i + 1:] if x.strip()), "")
        if following.startswith("|"):
            paragraph.paragraph_format.keep_with_next = True
        i += 1
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
